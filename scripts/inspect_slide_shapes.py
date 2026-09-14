import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\HD_EXTERNO\Moinho\Dados\Dados_analise_diagnostico\2026-09 Alianzo ENTHUS - Proposta Consultoria Comercial - Moinho Sete Irmaos V4.pptx')

print(f"Dimensões do Slide: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} polegadas")

slides_interesse = [6, 7, 8, 9, 11, 12, 14, 17, 18, 19, 20, 22, 28, 33, 34, 35]

for s_idx in slides_interesse:
    s = prs.slides[s_idx - 1]
    print(f"\n==================== SLIDE {s_idx} ====================")
    for sh_idx, shape in enumerate(s.shapes):
        t = ""
        if shape.has_text_frame:
            t = " ".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
        sh_type = shape.shape_type
        print(f"Shape {sh_idx}: {shape.name} | Type: {sh_type} | Pos: ({shape.left.inches:.2f}, {shape.top.inches:.2f}) | Dim: ({shape.width.inches:.2f}, {shape.height.inches:.2f}) | Text: {t[:60]}")
