import sys
import os
import shutil
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"E:\HD_EXTERNO\Moinho\Dados\Dados_analise_diagnostico\2026-09 Alianzo ENTHUS - Proposta Consultoria Comercial - Moinho Sete Irmaos V4.pptx"
backup_path = r"E:\HD_EXTERNO\Moinho\Dados\Dados_analise_diagnostico\2026-09 Alianzo ENTHUS - Proposta Consultoria Comercial - Moinho Sete Irmaos V4_backup.pptx"

# Backup de segurança
if not os.path.exists(backup_path):
    shutil.copyfile(pptx_path, backup_path)
    print(f"Backup criado em: {backup_path}")
else:
    print(f"Backup ja existe em: {backup_path}")

prs = Presentation(pptx_path)
print(f"\nTotal de Slides: {len(prs.slides)}")
print("="*80)

for idx, slide in enumerate(prs.slides):
    title = ""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
    
    if texts:
        title = texts[0]
    
    notes = ""
    if slide.has_notes_slide:
        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame:
            notes = notes_frame.text.strip()
            
    print(f"\n--- SLIDE {idx + 1} ---")
    print(f"Título / Primeiro Texto: {title}")
    print(f"Total blocos de texto: {len(texts)}")
    print(f"Conteúdo resumido: {' | '.join(texts[:4])}")
    if notes:
        print(f"Notas do Orador Existentes: {notes[:150]}...")
    else:
        print("Notas do Orador: (vazio)")
