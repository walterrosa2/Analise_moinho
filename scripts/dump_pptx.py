import sys
import os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\HD_EXTERNO\Moinho\Dados\Dados_analise_diagnostico\2026-09 Alianzo ENTHUS - Proposta Consultoria Comercial - Moinho Sete Irmaos V4.pptx')

with open("artifacts/analise_slides_pptx.txt", "w", encoding="utf-8") as out:
    for i, s in enumerate(prs.slides):
        txts = [p.text.strip() for shape in s.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs if p.text.strip()]
        out.write(f"\n=======================================================\n")
        out.write(f"SLIDE {i+1} (shapes: {len(s.shapes)})\n")
        out.write(f"=======================================================\n")
        for t in txts:
            out.write(f"  • {t}\n")
        
        # Verificar notas do orador
        if s.has_notes_slide and s.notes_slide.notes_text_frame:
            nt = s.notes_slide.notes_text_frame.text.strip()
            if nt:
                out.write(f"\n  [NOTAS DO ORADOR EXISTENTES]:\n  {nt}\n")

print(f"Salvo em artifacts/analise_slides_pptx.txt com {len(prs.slides)} slides.")
